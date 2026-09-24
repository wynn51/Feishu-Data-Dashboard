import argparse
import gc
import hashlib
import os
import re
import subprocess
import tempfile
import time
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_MONTH = "2026-06"
SITE_BATCHES = ("菲律宾", "马来", "泰国", "印尼1", "印尼2")
TEMPLATE_NAMES = {
    "菲律宾": "@净需求计算表-菲律宾.xlsx",
    "马来": "@净需求计算表-马来.xlsx",
    "泰国": "@净需求计算表-泰国.xlsx",
    "印尼1": "@净需求计算表-印尼1仓.xlsx",
    "印尼2": "@净需求计算表-印尼2仓.xlsx",
}
SHEET_ORDER = (
    "净需求表", "计划目标表", "货值情况", "千易库存数据", "千易入库单_收货中",
    "销量占比", "SKU信息表", "仓库信息表", "入库单信息", "在途批次数据",
)
SOURCE_SHEETS = {
    "千易库存数据", "千易入库单_收货中", "SKU信息表", "仓库信息表", "入库单信息", "在途批次数据"
}
FORMULA_ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#N/A"}
DEMAND_MONTH_GROUP_STARTS = (16, 22, 65, 71, 77, 83, 89, 95, 101, 107, 113)
HELPER_CLASS_TARGET_PREFIX = "类目目标"
HELPER_SKU_TARGET_PREFIX = "SKU目标"
HELPER_TARGET_HEADER_PATTERN = re.compile(r"^(类目目标|SKU目标)\d{4}\.\d{1,2}月$")


def target_months(month) :
    return [add_months(month, offset) for offset in range(1, 7)]


def log(message) :
    print(f"[{datetime.now():%H:%M:%S}] {message}", flush=True)


def verify_file_released(path) :
    if not path.exists():
        return
    probe = path.with_name(f".{path.name}.lockcheck.{os.getpid()}.tmp")
    try:
        path.replace(probe)
        probe.replace(path)
    except PermissionError as exc:
        raise PermissionError(f"文件仍被 Python、Excel 或其他程序占用，无法移动：{path}") from exc
    finally:
        if probe.exists() and not path.exists():
            probe.replace(path)


def release_resources_and_verify(paths) :
    gc.collect()
    time.sleep(0.5)
    for path in paths:
        verify_file_released(path)
    log("资源释放校验完成：输出 Excel 未被当前任务占用")


def normalize(value) :
    return "" if value is None else str(value).replace("\xa0", "").strip()


def parse_month(value) :
    try:
        return datetime.strptime(value, "%Y-%m").date().replace(day=1)
    except ValueError as exc:
        raise ValueError("执行月份必须使用 YYYY-MM，例如 2026-06") from exc


def next_month_first_day(month) :
    return (month.replace(day=28) + timedelta(days=4)).replace(day=1)


def add_months(month, offset) :
    index = month.year * 12 + month.month - 1 + offset
    return date(index // 12, index % 12 + 1, 1)


def month_label(month, suffix = "") :
    return f"{month.year}.{month.month}月{suffix}"


def demand_header_row(ws) :
    for row in range(1, min(ws.max_row, 5) + 1):
        if normalize(ws.cell(row, 16).value) == "类目目标":
            return row + 1
    raise ValueError("净需求表未找到类目目标表头")


def plan_target_columns(ws) :
    pattern = re.compile(r"^(\d{2}|\d{4})年(\d{1,2})月目标$")
    result = {}
    duplicates = []
    for column in range(1, ws.max_column + 1):
        match = pattern.match(normalize(ws.cell(1, column).value))
        if not match:
            continue
        year = int(match.group(1))
        if year < 100:
            year += 2000
        key = (year, int(match.group(2)))
        if key in result:
            duplicates.append(key)
        result[key] = column
    if duplicates:
        raise ValueError(f"计划目标表存在重复目标月份：{duplicates}")
    return result


def update_demand_month_window(wb, month) :
    demand = wb["净需求表"]
    targets = plan_target_columns(wb["计划目标表"])
    months = target_months(month)
    helper_month = month
    missing = [month_label(value) for value in months if (value.year, value.month) not in targets]
    if (helper_month.year, helper_month.month) not in targets:
        missing.append(month_label(helper_month))
    if missing:
        log("计划目标表缺少以下目标月份，对应目标列将按 0 写入并继续执行：" + "、".join(missing))

    header_row = demand_header_row(demand)
    for start in DEMAND_MONTH_GROUP_STARTS:
        for offset, value in enumerate(months):
            demand.cell(header_row, start + offset).value = month_label(value)
    for offset in range(4):
        demand.cell(header_row, 29 + offset).value = month_label(add_months(month, offset), "库存")

    formula_count = 0
    for offset, value in enumerate(months):
        demand_column = 16 + offset
        target_column = targets.get((value.year, value.month))
        target_letter = get_column_letter(target_column) if target_column else None
        lookup_index = target_column - 2 if target_column else None
        for row in range(header_row + 1, demand.max_row + 1):
            cell = demand.cell(row, demand_column)
            if not isinstance(cell.value, str) or not cell.value.startswith("="):
                continue
            if target_column is None:
                cell.value = 0
            elif "VLOOKUP" in cell.value:
                cell.value = f'=IFERROR(VLOOKUP(B{row},计划目标表!C:{target_letter},{lookup_index},0),0)'
            else:
                cell.value = f'=IFERROR(_xlfn.XLOOKUP(B{row},计划目标表!C:C,计划目标表!{target_letter}:{target_letter}),0)'
            formula_count += 1
    removed_helper_columns = remove_helper_target_columns(demand, header_row)
    if removed_helper_columns:
        log(f"净需求表已清理历史辅助目标列：{removed_helper_columns}")
    helper_class_col, helper_sku_col = ensure_helper_target_columns(demand, wb["计划目标表"], month, targets)
    helper_target_column = targets.get((month.year, month.month))
    if helper_target_column is None:
        log(f"计划目标表缺少辅助目标月份 {month_label(month)}，辅助类目目标按 0 写入并继续执行")
    formula_count += fill_helper_target_formulas(demand, header_row, helper_class_col, helper_sku_col, helper_target_column)
    return formula_count


def validate_demand_month_window(wb, month) :
    demand = wb["净需求表"]
    targets = plan_target_columns(wb["计划目标表"])
    months = target_months(month)
    header_row = demand_header_row(demand)
    for start in DEMAND_MONTH_GROUP_STARTS:
        actual = [normalize(demand.cell(header_row, start + offset).value) for offset in range(6)]
        expected = [month_label(value) for value in months]
        if actual != expected:
            raise ValueError(f"净需求表月份窗口错误：{actual} != {expected}")
    inventory_actual = [normalize(demand.cell(header_row, 29 + offset).value) for offset in range(4)]
    inventory_expected = [month_label(add_months(month, offset), "库存") for offset in range(4)]
    if inventory_actual != inventory_expected:
            raise ValueError(f"净需求表库存月份窗口错误：{inventory_actual} != {inventory_expected}")
    for offset, value in enumerate(months):
        target_column = targets.get((value.year, value.month))
        if target_column is None:
            bad_zero = [
                demand.cell(row, 16 + offset).coordinate
                for row in range(header_row + 1, demand.max_row + 1)
                if normalize(demand.cell(row, 6).value)
                and normalize(demand.cell(row, 16 + offset).value) not in {"0", "0.0"}
            ]
            if bad_zero:
                raise ValueError(f"净需求表 {get_column_letter(16 + offset)} 列缺少计划目标月份时未按 0 写入，样例：{bad_zero[:5]}")
        else:
            target_letter = get_column_letter(target_column)
            formulas = [
                demand.cell(row, 16 + offset).value
                for row in range(header_row + 1, demand.max_row + 1)
                if isinstance(demand.cell(row, 16 + offset).value, str)
                and demand.cell(row, 16 + offset).value.startswith("=")
            ]
            if not formulas or any(
                f"计划目标表!C:{target_letter}" not in formula
                and f"计划目标表!{target_letter}:{target_letter}" not in formula
                for formula in formulas
            ):
                raise ValueError(f"净需求表 {get_column_letter(16 + offset)} 列未完整引用 {month_label(value)} 目标列 {target_letter}")
    helper_class_col, helper_sku_col = helper_target_columns(demand, month, required=True)
    helper_target_column = targets.get((month.year, month.month))
    helper_formulas = [
        normalize(demand.cell(row, helper_class_col).value)
        for row in range(header_row + 1, demand.max_row + 1)
        if normalize(demand.cell(row, 6).value)
    ]
    if helper_target_column is None:
        if any(value not in {"0", "0.0"} for value in helper_formulas):
            raise ValueError(f"净需求表辅助列 {get_column_letter(helper_class_col)} 缺少计划目标月份时未按 0 写入")
    else:
        helper_target_letter = get_column_letter(helper_target_column)
        if not helper_formulas or any(f"计划目标表!{helper_target_letter}:{helper_target_letter}" not in formula for formula in helper_formulas if formula.startswith("=")):
            raise ValueError(f"净需求表辅助列 {get_column_letter(helper_class_col)} 未完整引用 {month_label(month)} 目标列 {helper_target_letter}")
    sku_formulas = [
        normalize(demand.cell(row, helper_sku_col).value)
        for row in range(header_row + 1, demand.max_row + 1)
        if normalize(demand.cell(row, 6).value)
    ]
    if not sku_formulas or any(
        f"{get_column_letter(helper_class_col)}{row}" not in normalize(demand.cell(row, helper_sku_col).value)
        for row in range(header_row + 1, demand.max_row + 1)
        if normalize(demand.cell(row, 6).value)
    ):
        raise ValueError(f"净需求表辅助列 {get_column_letter(helper_sku_col)} 未完整引用 {get_column_letter(helper_class_col)} 列")


def helper_target_labels(month) :
    label = month_label(month)
    return f"{HELPER_CLASS_TARGET_PREFIX}{label}", f"{HELPER_SKU_TARGET_PREFIX}{label}"


def remove_helper_target_columns(ws, header_row) :
    removed = []
    for column in range(ws.max_column, 0, -1):
        value = normalize(ws.cell(header_row, column).value)
        if HELPER_TARGET_HEADER_PATTERN.match(value):
            removed.append(f"{get_column_letter(column)}:{value}")
            ws.delete_cols(column, 1)
    return list(reversed(removed))


def helper_target_columns(ws, month, required = False) :
    header_row = demand_header_row(ws)
    class_label, sku_label = helper_target_labels(month)
    positions = {}
    for column in range(1, ws.max_column + 1):
        value = normalize(ws.cell(header_row, column).value)
        if value in {class_label, sku_label}:
            positions[value] = column
    if required and (class_label not in positions or sku_label not in positions):
        raise ValueError(f"净需求表缺少辅助目标列：{class_label}、{sku_label}")
    return positions.get(class_label), positions.get(sku_label)


def copy_column_style(ws, source_col, target_col) :
    ws.column_dimensions[get_column_letter(target_col)].width = ws.column_dimensions[get_column_letter(source_col)].width
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row, source_col)
        target = ws.cell(row, target_col)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def ensure_helper_target_columns(ws, target_ws, month, targets) :
    header_row = demand_header_row(ws)
    class_col, sku_col = helper_target_columns(ws, month)
    if class_col and sku_col:
        return class_col, sku_col
    class_col = ws.max_column + 1
    sku_col = class_col + 1
    if class_col > 1:
        copy_column_style(ws, class_col - 1, class_col)
        copy_column_style(ws, class_col - 1, sku_col)
    if header_row > 1:
        ws.cell(header_row - 1, class_col).value = "辅助备用列"
    class_label, sku_label = helper_target_labels(month)
    ws.cell(header_row, class_col).value = class_label
    ws.cell(header_row, sku_col).value = sku_label
    return class_col, sku_col


def fill_helper_target_formulas(ws, header_row, class_col, sku_col, target_column) :
    target_letter = get_column_letter(target_column) if target_column else None
    class_letter = get_column_letter(class_col)
    count = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if not normalize(ws.cell(row, 6).value):
            continue
        if target_letter is None:
            ws.cell(row, class_col).value = 0
        else:
            ws.cell(row, class_col).value = f'=IFERROR(_xlfn.XLOOKUP(B{row},计划目标表!C:C,计划目标表!{target_letter}:{target_letter}),0)'
        ws.cell(row, sku_col).value = f'={class_letter}{row}*J{row}*K{row}'
        count += 2
    return count


def sync_dq_in_transit_formula(wb) :
    ws = wb["净需求表"]
    header_row = demand_header_row(ws)
    changed = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if not normalize(ws.cell(row, 6).value):
            continue
        formula = (
            f'=IFERROR(SUMPRODUCT(--(EXACT(TRIM(CLEAN(入库单信息!C:C)),TRIM(CLEAN($F{row})))),'
            f'--(NOT(EXACT(TRIM(CLEAN(入库单信息!A:A)),"A未发"))),入库单信息!D:D),0)+M{row}'
        )
        cell = ws.cell(row, 121)
        if cell.value != formula:
            cell.value = formula
            changed += 1
    return changed


def file_hash(path) :
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_map(ws, row = 1) :
    return {
        normalize(ws.cell(row, column).value): column
        for column in range(1, ws.max_column + 1)
        if normalize(ws.cell(row, column).value)
    }


def require_columns(ws, names, row = 1) :
    mapping = header_map(ws, row)
    missing = [name for name in names if name not in mapping]
    if missing:
        raise KeyError(f"{ws.title} 缺少字段：{missing}")
    return mapping


def valid_row_count(ws, key_column, first_row) :
    return sum(1 for row in range(first_row, ws.max_row + 1) if normalize(ws.cell(row, key_column).value))


def copy_style(source, target) :
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    target.number_format = source.number_format


def capture_row_style(ws, row, max_column) :
    return [copy(ws.cell(row, column)._style) for column in range(1, max_column + 1)], ws.row_dimensions[row].height


def apply_row_style(ws, row, captured) :
    styles, height = captured
    for column, style in enumerate(styles, 1):
        ws.cell(row, column)._style = copy(style)
    ws.row_dimensions[row].height = height


def clear_data_rows(ws) :
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)


def replace_sheet_with_values(target_ws, source_ws) :
    for merged in list(target_ws.merged_cells.ranges):
        target_ws.unmerge_cells(str(merged))
    old_rows, old_columns = target_ws.max_row, target_ws.max_column
    if old_columns:
        target_ws.delete_cols(1, old_columns)
    if old_rows:
        target_ws.delete_rows(1, old_rows)
    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = target_ws[source_cell.coordinate]
            target_cell.value = source_cell.value
            copy_style(source_cell, target_cell)
    for key, dimension in source_ws.column_dimensions.items():
        target_ws.column_dimensions[key].width = dimension.width
        target_ws.column_dimensions[key].hidden = dimension.hidden
    for index, dimension in source_ws.row_dimensions.items():
        target_ws.row_dimensions[index].height = dimension.height
        target_ws.row_dimensions[index].hidden = dimension.hidden
    for merged in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged))
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.auto_filter.ref = source_ws.auto_filter.ref
    return valid_row_count(target_ws, 1, 2)


def translate_formula(prototype, target_coordinate) :
    formula, origin = prototype
    return Translator(formula, origin=origin).translate_formula(target_coordinate)


def resolve_paths(root, month, site_batch) :
    snapshot = next_month_first_day(month).strftime("%Y-%m-%d")
    month_cn = f"{month.year:04d}年{month.month:02d}月"
    template = root / "源数据" / "净需求计算表" / TEMPLATE_NAMES[site_batch]
    paths = {
        "template": template,
        "formula_blueprint": root / "源数据" / "净需求计算表" / TEMPLATE_NAMES["菲律宾"],
        "mapping": root / "源数据" / "映射表" / "SKU、店铺、仓库映射表.xlsx",
        "stock": root / "源数据" / "千易库存数据" / f"库存查询_{snapshot}.xlsx",
        "receiving": root / "源数据" / "千易入库单数据" / f"入库单明细_收货中_{snapshot}.xlsx",
        "batch": root / "数据中间表" / "在途批次表" / f"在途批次表_{month_cn}.xlsx",
        "output": root / "数据中间表" / "净需求计算表" / f"净需求计算表_{site_batch}_{month_cn}.xlsx",
    }
    missing = [str(path) for name, path in paths.items() if name != "output" and not path.exists()]
    if missing:
        raise FileNotFoundError("缺少以下输入文件：\n" + "\n".join(missing))
    return paths


def warehouse_site_map(mapping_ws) :
    columns = require_columns(mapping_ws, ["仓库代码", "对应站点"])
    return {
        normalize(mapping_ws.cell(row, columns["仓库代码"]).value):
        normalize(mapping_ws.cell(row, columns["对应站点"]).value)
        for row in range(2, mapping_ws.max_row + 1)
        if normalize(mapping_ws.cell(row, columns["仓库代码"]).value)
    }


def station_matches(site_batch, warehouse, warehouse_sites) :
    if site_batch == "印尼1":
        return "印尼云仓1" in warehouse
    if site_batch == "印尼2":
        return "印尼云仓2" in warehouse
    return warehouse_sites.get(warehouse) == site_batch


def select_formula_prototype(target_ws, fallback_ws, column) :
    target = target_ws.cell(2, column)
    if isinstance(target.value, str) and target.value.startswith("="):
        return target.value, target.coordinate
    fallback = fallback_ws.cell(2, column)
    if isinstance(fallback.value, str) and fallback.value.startswith("="):
        return fallback.value, fallback.coordinate
    raise ValueError(f"缺少公式原型：{target_ws.title}!{get_column_letter(column)}2")


def write_stock_sheet(target_ws, fallback_ws, source_ws, site_batch, warehouse_sites) :
    expected_headers = [
        "商品编码", "商品名称", "主SKU分类", "仓库", "站点（中文）", "可用量", "在途量", "未收到库存",
        "实际在途量", "库存合计", "3天销量", "7天销量", "15天销量", "30天销量", "60天销量", "90天销量",
    ]
    if [normalize(target_ws.cell(1, c).value) for c in range(1, 17)] != expected_headers:
        raise ValueError("模板千易库存数据表头与预期不一致")
    columns = require_columns(
        source_ws,
        ["商品编码", "商品名称", "仓库", "可用量", "在途量", "3天销量", "7天销量", "15天销量", "30天销量", "60天销量", "90天销量"],
        row=2,
    )
    formula_columns = (3, 5, 8, 9, 10)
    prototypes = {column: select_formula_prototype(target_ws, fallback_ws, column) for column in formula_columns}
    style_ws = target_ws if target_ws.max_row >= 2 else fallback_ws
    row_style = capture_row_style(style_ws, 2, 16)
    rows = []
    for source_row in range(3, source_ws.max_row + 1):
        sku = normalize(source_ws.cell(source_row, columns["商品编码"]).value)
        warehouse = normalize(source_ws.cell(source_row, columns["仓库"]).value)
        if not sku or not station_matches(site_batch, warehouse, warehouse_sites):
            continue
        rows.append([
            source_ws.cell(source_row, columns["商品编码"]).value,
            source_ws.cell(source_row, columns["商品名称"]).value,
            None,
            source_ws.cell(source_row, columns["仓库"]).value,
            None,
            source_ws.cell(source_row, columns["可用量"]).value,
            source_ws.cell(source_row, columns["在途量"]).value,
            None, None, None,
            source_ws.cell(source_row, columns["3天销量"]).value,
            source_ws.cell(source_row, columns["7天销量"]).value,
            source_ws.cell(source_row, columns["15天销量"]).value,
            source_ws.cell(source_row, columns["30天销量"]).value,
            source_ws.cell(source_row, columns["60天销量"]).value,
            source_ws.cell(source_row, columns["90天销量"]).value,
        ])
    clear_data_rows(target_ws)
    for target_row, values in enumerate(rows, 2):
        apply_row_style(target_ws, target_row, row_style)
        for column, value in enumerate(values, 1):
            target_ws.cell(target_row, column).value = value
        for column, prototype in prototypes.items():
            target_ws.cell(target_row, column).value = translate_formula(prototype, target_ws.cell(target_row, column).coordinate)
    log(f"Sheet 完成：千易库存数据，输入有效行={valid_row_count(source_ws, columns['商品编码'], 3)}，站点输出行={len(rows)}，公式={len(rows) * 5}")
    return len(rows)


def write_receiving_sheet(target_ws, fallback_ws, source_ws, site_batch, warehouse_sites) :
    columns = require_columns(source_ws, ["入库单号", "收货仓", "商品编码", "预计数量", "仓库上架数量"], row=1)
    target_columns = require_columns(target_ws, ["批注", "预计数量", "仓库上架数量", "差异数据"], row=1)
    formula_columns = (8,)
    prototypes = {column: select_formula_prototype(target_ws, fallback_ws, column) for column in formula_columns}
    style_ws = target_ws if target_ws.max_row >= 2 else fallback_ws
    row_style = capture_row_style(style_ws, 2, 48)
    remark_letter = get_column_letter(target_columns["批注"])
    expected_letter = get_column_letter(target_columns["预计数量"])
    shelved_letter = get_column_letter(target_columns["仓库上架数量"])
    difference_column = target_columns["差异数据"]
    rows = []
    for source_row in range(2, source_ws.max_row + 1):
        order = normalize(source_ws.cell(source_row, columns["入库单号"]).value)
        warehouse = normalize(source_ws.cell(source_row, columns["收货仓"]).value)
        if not order or not station_matches(site_batch, warehouse, warehouse_sites):
            continue
        source_values = [source_ws.cell(source_row, column).value for column in range(1, 47)]
        rows.append(source_values[:7] + [None] + source_values[7:32] + [None] + source_values[32:])
    clear_data_rows(target_ws)
    for target_row, values in enumerate(rows, 2):
        apply_row_style(target_ws, target_row, row_style)
        for column, value in enumerate(values, 1):
            target_ws.cell(target_row, column).value = value
        for column, prototype in prototypes.items():
            target_ws.cell(target_row, column).value = translate_formula(prototype, target_ws.cell(target_row, column).coordinate)
        target_ws.cell(
            target_row,
            difference_column,
            f'=IF(ISNUMBER(FIND("点完",{remark_letter}{target_row})),MAX({expected_letter}{target_row}-{shelved_letter}{target_row},0),0)',
        )
    log(f"Sheet 完成：千易入库单_收货中，输入有效行={valid_row_count(source_ws, columns['入库单号'], 2)}，站点输出行={len(rows)}，公式={len(rows) * 2}")
    return len(rows)


def write_inbound_sheet(target_ws, source_ws, site_batch) :
    expected_headers = ["入库单号", "仓库", "商品编码", "预计数量", "交货时间", "到货时间"]
    if [normalize(target_ws.cell(1, column).value) for column in range(1, 7)] != expected_headers:
        raise ValueError(f"模板入库单信息表头与预期不一致：{expected_headers}")
    columns = require_columns(
        source_ws,
        ["入库单号", "仓库", "站点_批次", "商品编码", "预计数量Value", "发货时间", "到仓时间整合"],
    )
    style_row = 2 if target_ws.max_row >= 2 else 1
    row_style = capture_row_style(target_ws, style_row, 6)
    rows = []
    for source_row in range(2, source_ws.max_row + 1):
        if normalize(source_ws.cell(source_row, columns["站点_批次"]).value) != site_batch:
            continue
        rows.append([
            source_ws.cell(source_row, columns["入库单号"]).value,
            source_ws.cell(source_row, columns["仓库"]).value,
            source_ws.cell(source_row, columns["商品编码"]).value,
            source_ws.cell(source_row, columns["预计数量Value"]).value,
            source_ws.cell(source_row, columns["发货时间"]).value,
            source_ws.cell(source_row, columns["到仓时间整合"]).value,
        ])
    clear_data_rows(target_ws)
    for target_row, values in enumerate(rows, 2):
        apply_row_style(target_ws, target_row, row_style)
        for column, value in enumerate(values, 1):
            target_ws.cell(target_row, column).value = value
    log(f"Sheet 完成：入库单信息，旧数据已清空，站点_批次={site_batch}，输出行={len(rows)}，静态值写入")
    return len(rows)


def write_batch_sheet(target_ws, source_ws) :
    expected = [
        "商品编码", "第一批入库单号", "第一批数量", "第一批交货时间", "第一批预计到货时间",
        "第二批入库单号", "第二批数量", "第二批交货时间", "第二批预计到货时间",
        "第三批入库单号", "第三批数量", "第三批交货时间", "第三批预计到货时间",
        "第四批入库单号", "第四批数量", "第四批交货时间", "第四批预计到货时间",
    ]
    if [normalize(source_ws.cell(1, c).value) for c in range(1, 18)] != expected:
        raise ValueError(f"{source_ws.title} 表头与预期不一致")
    style_row = 2 if target_ws.max_row >= 2 else 1
    row_style = capture_row_style(target_ws, style_row, 17)
    rows = [
        [source_ws.cell(row, column).value for column in range(1, 18)]
        for row in range(2, source_ws.max_row + 1)
        if normalize(source_ws.cell(row, 1).value)
    ]
    clear_data_rows(target_ws)
    for target_row, values in enumerate(rows, 2):
        apply_row_style(target_ws, target_row, row_style)
        for column, value in enumerate(values, 1):
            target_ws.cell(target_row, column).value = value
    log(f"Sheet 完成：在途批次数据，来源={source_ws.title}，输出SKU行={len(rows)}，静态值写入")
    return len(rows)


def internal_formula_map(wb) :
    formulas = {}
    for ws in wb.worksheets:
        if ws.title in SOURCE_SHEETS:
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[(ws.title, cell.coordinate)] = cell.value
    return formulas


def error_map(path) :
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        result = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                        result[(ws.title, cell.coordinate)] = cell.value
        return result
    finally:
        wb.close()


def recalculate_with_excel(path, timeout_seconds = 1200) :
    log("开始调用 Microsoft Excel 全量计算并刷新公式缓存")
    script = r'''
param([string]$WorkbookPath, [int]$TimeoutSeconds)
$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.AskToUpdateLinks = $false
    $book = $excel.Workbooks.Open($WorkbookPath, 0, $false)
    $excel.CalculateFullRebuild()
    $start = Get-Date
    while ($excel.CalculationState -ne 0) {
        if (((Get-Date) - $start).TotalSeconds -gt $TimeoutSeconds) { throw 'Excel 计算超时' }
        Start-Sleep -Milliseconds 500
    }
    $book.Save()
}
finally {
    if ($book -ne $null) { $book.Close($true) | Out-Null; [void][Runtime.InteropServices.Marshal]::ReleaseComObject($book) }
    if ($excel -ne $null) { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel) }
    [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}
'''
    started = time.monotonic()
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as handle:
            handle.write(script)
            script_path = Path(handle.name)
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script_path),
             "-WorkbookPath", str(path.resolve()), "-TimeoutSeconds", str(timeout_seconds)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout_seconds + 60,
        )
        if completed.returncode:
            raise RuntimeError("Excel 公式缓存刷新失败：\n" + completed.stderr.strip())
    finally:
        if script_path:
            try:
                script_path.unlink()
            except FileNotFoundError:
                pass
    elapsed = time.monotonic() - started
    log(f"Microsoft Excel 全量计算、保存和关闭完成，耗时={elapsed:.1f}秒")
    return elapsed


def validate_output(
    path,
    template,
    template_formulas,
    expected_rows,
    month,
) :
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    try:
        if tuple(formula_wb.sheetnames) != SHEET_ORDER:
            raise ValueError(f"输出 Sheet 顺序不正确：{formula_wb.sheetnames}")
        validate_demand_month_window(formula_wb, month)
        if internal_formula_map(formula_wb) != template_formulas:
            raise ValueError("业务计算 Sheet 的公式内容与模板不一致，已停止输出")
        if getattr(formula_wb, "_external_links", []):
            raise ValueError("输出工作簿存在外部工作簿链接")
        actual_rows = {
            "千易库存数据": valid_row_count(formula_wb["千易库存数据"], 1, 2),
            "千易入库单_收货中": valid_row_count(formula_wb["千易入库单_收货中"], 1, 2),
            "入库单信息": valid_row_count(formula_wb["入库单信息"], 1, 2),
            "在途批次数据": valid_row_count(formula_wb["在途批次数据"], 1, 2),
        }
        if actual_rows != expected_rows:
            raise ValueError(f"来源 Sheet 行数校验失败：预期={expected_rows}，实际={actual_rows}")
        inbound_ws = formula_wb["入库单信息"]
        inbound_headers = [normalize(inbound_ws.cell(1, column).value) for column in range(1, 7)]
        if inbound_headers != ["入库单号", "仓库", "商品编码", "预计数量", "交货时间", "到货时间"]:
            raise ValueError(f"入库单信息字段顺序不正确：{inbound_headers}")
        blank_warehouses = [
            row for row in range(2, inbound_ws.max_row + 1)
            if normalize(inbound_ws.cell(row, 1).value) and not normalize(inbound_ws.cell(row, 2).value)
        ]
        if blank_warehouses:
            raise ValueError(f"入库单信息存在空仓库，示例行：{blank_warehouses[:20]}")
        stock_formulas = sum(
            1 for row in range(2, formula_wb["千易库存数据"].max_row + 1)
            for column in (3, 5, 8, 9, 10)
            if isinstance(formula_wb["千易库存数据"].cell(row, column).value, str)
            and formula_wb["千易库存数据"].cell(row, column).value.startswith("=")
        )
        receiving_formulas = sum(
            1 for row in range(2, formula_wb["千易入库单_收货中"].max_row + 1)
            for column in (8, 34)
            if isinstance(formula_wb["千易入库单_收货中"].cell(row, column).value, str)
            and formula_wb["千易入库单_收货中"].cell(row, column).value.startswith("=")
        )
        if stock_formulas != expected_rows["千易库存数据"] * 5:
            raise ValueError("千易库存数据新增行公式数量不正确")
        if receiving_formulas != expected_rows["千易入库单_收货中"] * 2:
            raise ValueError("千易入库单_收货中新增行公式数量不正确")
    finally:
        formula_wb.close()

    baseline = error_map(template)
    current = error_map(path)
    unexpected = {
        coordinate: value
        for coordinate, value in current.items()
        if value != "#N/A"
        and baseline.get(coordinate) != value
        and coordinate not in template_formulas
    }
    if unexpected:
        sample = [f"{sheet}!{cell}={value}" for (sheet, cell), value in list(unexpected.items())[:30]]
        raise ValueError("发现模板基线之外的公式错误：\n" + "\n".join(sample))
    counts = Counter(current.values())
    log(
        "结果校验完成："
        f"来源行数={expected_rows}，错误缓存统计={dict(counts)}；"
        "#N/A 按既定规则保留，模板原有错误已识别且未新增"
    )


def build_workbook(paths, site_batch, month) :
    template_hash_before = file_hash(paths["template"])
    template_wb = load_workbook(paths["template"], data_only=False, read_only=False)
    blueprint_wb = load_workbook(paths["formula_blueprint"], data_only=False, read_only=False)
    mapping_values = load_workbook(paths["mapping"], data_only=True, read_only=False)
    stock_values = load_workbook(paths["stock"], data_only=True, read_only=False)
    receiving_values = load_workbook(paths["receiving"], data_only=True, read_only=False)
    batch_values = load_workbook(paths["batch"], data_only=True, read_only=False)
    try:
        if tuple(template_wb.sheetnames) != SHEET_ORDER:
            raise ValueError(f"模板 Sheet 顺序不正确：{template_wb.sheetnames}")
        month_formula_count = update_demand_month_window(template_wb, month)
        dq_formula_count = sync_dq_in_transit_formula(template_wb)
        log(f"Sheet 完成：净需求表主月份窗口={month_label(add_months(month, 1))}至{month_label(add_months(month, 6))}，辅助目标月份={month_label(month)}，计划目标引用公式={month_formula_count}")
        log(f"Sheet 完成：净需求表 DQ列在途公式校准，更新公式数={dq_formula_count}")
        template_formulas = internal_formula_map(template_wb)
        sku_rows = replace_sheet_with_values(template_wb["SKU信息表"], mapping_values["SKU信息表"])
        log(f"Sheet 完成：SKU信息表，来源={paths['mapping']} / SKU信息表，输出行={sku_rows}")
        warehouse_rows = replace_sheet_with_values(template_wb["仓库信息表"], mapping_values["仓库信息表"])
        log(f"Sheet 完成：仓库信息表，来源={paths['mapping']} / 仓库信息表，输出行={warehouse_rows}")
        sites = warehouse_site_map(mapping_values["仓库信息表"])
        stock_rows = write_stock_sheet(
            template_wb["千易库存数据"], blueprint_wb["千易库存数据"],
            stock_values[stock_values.sheetnames[0]], site_batch, sites,
        )
        receiving_rows = write_receiving_sheet(
            template_wb["千易入库单_收货中"], blueprint_wb["千易入库单_收货中"],
            receiving_values[receiving_values.sheetnames[0]], site_batch, sites,
        )
        order_sheet = batch_values["入库单表"]
        inbound_rows = write_inbound_sheet(template_wb["入库单信息"], order_sheet, site_batch)
        batch_sheet_name = f"SKU批次表_{site_batch}"
        if batch_sheet_name not in batch_values.sheetnames:
            raise KeyError(f"在途批次表缺少 Sheet：{batch_sheet_name}")
        batch_rows = write_batch_sheet(template_wb["在途批次数据"], batch_values[batch_sheet_name])

        template_wb.calculation.fullCalcOnLoad = True
        template_wb.calculation.forceFullCalc = True
        template_wb.calculation.calcMode = "auto"
        template_wb.save(paths["building"])
        expected_rows = {
            "千易库存数据": stock_rows,
            "千易入库单_收货中": receiving_rows,
            "入库单信息": inbound_rows,
            "在途批次数据": batch_rows,
        }
    finally:
        template_wb.close()
        blueprint_wb.close()
        mapping_values.close()
        stock_values.close()
        receiving_values.close()
        batch_values.close()
    if file_hash(paths["template"]) != template_hash_before:
        raise RuntimeError("源模板哈希发生变化，已停止执行")
    log(f"临时工作簿写入完成：{paths['building']}；源模板未修改")
    return expected_rows, template_formulas


def run(root, month_text, site_batch, output_dir = None) :
    if site_batch not in SITE_BATCHES:
        raise ValueError(f"站点批次必须是：{'、'.join(SITE_BATCHES)}")
    root = root.resolve()
    month = parse_month(month_text)
    paths = resolve_paths(root, month, site_batch)
    if output_dir:
        output_dir = Path(output_dir).expanduser()
        if not output_dir.is_absolute():
            output_dir = root / output_dir
        month_cn = f"{month.year:04d}年{month.month:02d}月"
        paths["output"] = output_dir.resolve() / f"净需求计算表_{site_batch}_{month_cn}.xlsx"
    temp_dir = root / ".codex_tmp"
    temp_dir.mkdir(exist_ok=True)
    paths["building"] = temp_dir / f"{paths['output'].stem}_building.xlsx"
    paths["output"].parent.mkdir(parents=True, exist_ok=True)

    log(f"根目录：{root}")
    log(f"执行月份：{month_text}；站点批次：{site_batch}")
    log(f"库存及收货中快照日期：{next_month_first_day(month):%Y-%m-%d}；不检索或回退其他日期")
    for name in ("template", "mapping", "stock", "receiving", "batch", "output"):
        log(f"路径确认 {name}：{paths[name]}")

    try:
        expected_rows, template_formulas = build_workbook(paths, site_batch, month)
        recalculate_with_excel(paths["building"])
        validate_output(paths["building"], paths["template"], template_formulas, expected_rows, month)
        os.replace(paths["building"], paths["output"])
        log(f"执行成功：{paths['output']}")
        return paths["output"]
    finally:
        try:
            paths["building"].unlink()
        except FileNotFoundError:
            pass
        release_resources_and_verify([paths["building"], paths["output"]])


def main() :
    parser = argparse.ArgumentParser(description="按月份和站点批次生成净需求计算表")
    parser.add_argument("--root", type=Path, default=ROOT_DIR, help="项目根目录")
    parser.add_argument("--month", default=DEFAULT_MONTH, help="执行月份，格式 YYYY-MM")
    parser.add_argument("--site", choices=SITE_BATCHES, help="指定单个站点批次；不传则依次生成全部站点")
    parser.add_argument("--output-dir", type=Path, help="输出目录；不传则保存到默认净需求计算表目录")
    args = parser.parse_args()
    sites = (args.site,) if args.site else SITE_BATCHES
    outputs = []
    for index, site in enumerate(sites, 1):
        log(f"开始执行站点 {index}/{len(sites)}：{site}")
        outputs.append(run(args.root, args.month, site, args.output_dir))
    log(f"全部站点执行完成：数量={len(outputs)}")


if __name__ == "__main__":
    main()
