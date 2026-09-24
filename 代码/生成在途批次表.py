import argparse
import gc
import os
import shutil
import subprocess
import tempfile
import time
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_MONTH = "2026-05"
SITE_BATCHES = ("泰国", "印尼1", "印尼2", "马来", "菲律宾")
UNSHIPPED_SOURCE_SHEETS = ("印尼1", "印尼2", "菲律宾", "马来", "泰国")
FINAL_SHEET_ORDER = (
    "千易入库单_新建", "A未发数据表", "入库单表", "SKU批次表_泰国",
    "SKU批次表_印尼1", "SKU批次表_印尼2", "SKU批次表_马来", "SKU批次表_菲律宾",
    "飞书-入库单跟进表", "SKU信息表", "店铺信息表", "仓库信息表",
)
SOURCE_REFRESH_SHEETS = (
    "千易入库单_新建", "A未发数据表", "飞书-入库单跟进表",
    "SKU信息表", "店铺信息表", "仓库信息表", "入库单表",
    "SKU批次表_泰国", "SKU批次表_印尼1", "SKU批次表_印尼2", "SKU批次表_马来", "SKU批次表_菲律宾",
)
BATCH_HEADERS = [
    "商品编码",
    "第一批入库单号", "第一批数量", "第一批交货时间", "第一批预计到货时间",
    "第二批入库单号", "第二批数量", "第二批交货时间", "第二批预计到货时间",
    "第三批入库单号", "第三批数量", "第三批交货时间", "第三批预计到货时间",
    "第四批入库单号", "第四批数量", "第四批交货时间", "第四批预计到货时间",
]
FORMULA_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Microsoft YaHei", size=10, color="000000")
THIN_SIDE = Side(style="thin", color="D9E2F3")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGNMENT = Alignment(vertical="center", wrap_text=False)
SHORT_DATE_FORMAT = "yyyy-mm-dd"


def log(message) :
    print(f"[{datetime.now():%H:%M:%S}] {message}", flush=True)


def verify_file_released(path) :
    if not path.exists():
        return
    probe = path.with_name(f".{path.name}.lockcheck.{os.getpid()}.tmp")
    try:
        path.replace(probe)
        probe.replace(path)
    except PermissionError as exc:
        raise PermissionError(f"文件仍被 Python、Excel 或其他程序占用，无法移动：{path}") from exc
    finally:
        if probe.exists() and not path.exists():
            probe.replace(path)


def release_resources_and_verify(paths) :
    gc.collect()
    time.sleep(0.5)
    for path in paths:
        verify_file_released(path)
    log("资源释放校验完成：输出 Excel 未被当前任务占用")


def normalize(value) :
    return "" if value is None else str(value).replace("\xa0", "").strip()


def header_map(ws, row = 1) :
    return {normalize(ws.cell(row, col).value): col for col in range(1, ws.max_column + 1) if normalize(ws.cell(row, col).value)}


def require_columns(ws, names, row = 1) :
    mapping = header_map(ws, row)
    missing = [name for name in names if name not in mapping]
    if missing:
        raise KeyError(f"{ws.title} 缺少字段：{missing}")
    return mapping


def parse_month(value) :
    try:
        return datetime.strptime(value, "%Y-%m").date().replace(day=1)
    except ValueError as exc:
        raise ValueError("月度日期必须使用 YYYY-MM，例如 2026-05") from exc


def next_month_first_day(month) :
    return (month.replace(day=28) + timedelta(days=4)).replace(day=1)


def resolve_paths(root, month) :
    snapshot_date = next_month_first_day(month)
    snapshot = snapshot_date.strftime("%Y-%m-%d")
    month_cn = f"{month.year:04d}年{month.month:02d}月"
    paths = {
        "target": root / "数据中间表" / "在途批次表" / f"在途批次表_{month_cn}.xlsx",
        "in_new": root / "源数据" / "千易入库单数据" / f"入库单明细_新建_{snapshot}.xlsx",
        "tracking": root / "源数据" / "飞书人工维护数据" / "入库单跟踪表" / f"入库单跟踪表_{snapshot}.xlsx",
        "stock": root / "源数据" / "千易库存数据" / f"库存查询_{snapshot}.xlsx",
        "unshipped": root / "源数据" / "飞书人工维护数据" / "A未发数据表" / f"A未发数据表_{month_cn}.xlsx",
        "mapping": root / "源数据" / "映射表" / "SKU、店铺、仓库映射表.xlsx",
    }
    missing = [str(path) for name, path in paths.items() if name != "target" and not path.exists()]
    if missing:
        raise FileNotFoundError("缺少以下文件：\n" + "\n".join(missing))
    return paths


def validate_stock_source(path) :
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        require_columns(ws, ["商品编码", "仓库", "可用量", "在途量"], row=2)
        valid_rows = sum(1 for row in range(3, ws.max_row + 1) if normalize(ws.cell(row, 1).value))
        if valid_rows == 0:
            raise ValueError(f"库存查询没有有效数据：{path}")
        return valid_rows
    finally:
        wb.close()


def apply_short_date_format(ws, columns, first_row = 2) :
    for row in range(first_row, ws.max_row + 1):
        for col in columns:
            ws.cell(row, col).number_format = SHORT_DATE_FORMAT


def is_short_date_format(number_format) :
    normalized = normalize(number_format).lower().replace("\\", "")
    return normalized in {"yyyy-mm-dd", "yyyy/m/d", "yyyy/mm/dd"}


def default_column_width(header) :
    if "入库单号" in header:
        return 26
    if "仓库" in header:
        return 28
    if "商品编码" in header or header == "SKU":
        return 20
    if "时间" in header or "日期" in header:
        return 18
    if "数量" in header or "天数" in header:
        return 14
    if "站点" in header:
        return 16
    return 18


def style_header(ws, headers) :
    ws.row_dimensions[1].height = 30
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col)
        if not cell.has_style:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGNMENT
            cell.border = CELL_BORDER
        letter = get_column_letter(col)
        if letter not in ws.column_dimensions or ws.column_dimensions[letter].width is None:
            ws.column_dimensions[letter].width = default_column_width(header)
    ws.freeze_panes = "A2"


def style_data_cell(cell) :
    if not cell.has_style:
        cell.font = DATA_FONT
        cell.alignment = DATA_ALIGNMENT
        cell.border = CELL_BORDER


def create_base_workbook() :
    wb = Workbook()
    wb.remove(wb.active)
    for name in FINAL_SHEET_ORDER:
        wb.create_sheet(name)
    return wb


def ensure_workbook_structure(wb) :
    for name in FINAL_SHEET_ORDER:
        if name not in wb.sheetnames:
            wb.create_sheet(name)


def load_or_create_workbook(target) :
    if target.exists():
        log(f"使用已执行过的在途批次表并重新刷新数据：{target}")
        wb = load_workbook(target, read_only=False, data_only=False)
        ensure_workbook_structure(wb)
        return wb
    target.parent.mkdir(parents=True, exist_ok=True)
    log(f"目标工作簿不存在，将直接从脚本内置结构、公式和样式初始化，不引用历史模板：{target}")
    wb = create_base_workbook()
    ensure_workbook_structure(wb)
    return wb


def copy_row_style(ws, source_row, target_row, max_col) :
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def reset_sheet(ws, headers) :
    old_max_row = ws.max_row
    if old_max_row > 1:
        ws.delete_rows(2, old_max_row - 1)
    for col, header in enumerate(headers, 1):
        ws.cell(1, col).value = header
    if ws.max_column > len(headers):
        ws.delete_cols(len(headers) + 1, ws.max_column - len(headers))
    style_header(ws, headers)


def clear_refresh_sheets(wb) :
    cleared = {}
    for name in SOURCE_REFRESH_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        old_rows = max(ws.max_row - 1, 0)
        if old_rows:
            ws.delete_rows(2, old_rows)
        cleared[name] = old_rows
    log(f"运行前清空源数据/派生源Sheet旧数据行：{cleared}")


def append_rows(ws, rows, style_row = 2) :
    count = 0
    for values in rows:
        target_row = ws.max_row + 1
        if target_row > 2 and ws.max_row >= style_row:
            copy_row_style(ws, style_row, target_row, len(values))
        for col, value in enumerate(values, 1):
            cell = ws.cell(target_row, col)
            cell.value = value
            style_data_cell(cell)
        count += 1
    return count


def copy_source_sheet(source_path, source_name, target_ws) :
    source_wb = load_workbook(source_path, read_only=False, data_only=False)
    try:
        source_ws = source_wb[source_name]
        if target_ws.max_row:
            target_ws.delete_rows(1, target_ws.max_row)
        if target_ws.max_column:
            target_ws.delete_cols(1, target_ws.max_column)
        for row in source_ws.iter_rows():
            for cell in row:
                target_cell = target_ws[cell.coordinate]
                target_cell.value = cell.value
                if cell.has_style:
                    target_cell.font = copy(cell.font)
                    target_cell.fill = copy(cell.fill)
                    target_cell.border = copy(cell.border)
                    target_cell.alignment = copy(cell.alignment)
                    target_cell.protection = copy(cell.protection)
                target_cell.number_format = cell.number_format
        for key, dim in source_ws.column_dimensions.items():
            target_ws.column_dimensions[key].width = dim.width
        target_ws.freeze_panes = source_ws.freeze_panes
        return max(source_ws.max_row - 1, 0)
    finally:
        source_wb.close()


def warehouse_site_map(mapping_path) :
    wb = load_workbook(mapping_path, read_only=False, data_only=True)
    try:
        ws = wb["仓库信息表"]
        cols = require_columns(ws, ["仓库代码", "对应站点"])
        result = {}
        for row in range(2, ws.max_row + 1):
            warehouse = normalize(ws.cell(row, cols["仓库代码"]).value)
            site = normalize(ws.cell(row, cols["对应站点"]).value)
            if warehouse:
                result[warehouse] = site
        return result
    finally:
        wb.close()


def station_batch(site, warehouse) :
    if site != "印尼":
        return site
    if "印尼云仓1" in warehouse:
        return "印尼1"
    if "印尼云仓2" in warehouse:
        return "印尼2"
    raise ValueError(f"印尼记录无法识别仓库粒度：{warehouse!r}")


def load_tracking(path) :
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        cols = require_columns(ws, ["仓库", "站点", "入库单号", "发货时间", "到仓时间", "到仓时间整合", "是否到货"])
        rows = []
        lookup = {}
        for row in range(2, ws.max_row + 1):
            order_id = normalize(ws.cell(row, cols["入库单号"]).value)
            if not order_id:
                continue
            values = [ws.cell(row, cols[name]).value for name in ["仓库", "站点", "入库单号", "发货时间", "到仓时间", "到仓时间整合", "是否到货"]]
            rows.append(values + [None])
            lookup[order_id] = (values[3], values[4], values[5])
        return rows, lookup
    finally:
        wb.close()


def load_new_inbound(path, sites, tracking) :
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        cols = require_columns(ws, ["入库单号", "收货仓", "商品编码", "预计数量"])
        rows = []
        for row in range(2, ws.max_row + 1):
            order_id = normalize(ws.cell(row, cols["入库单号"]).value)
            sku = normalize(ws.cell(row, cols["商品编码"]).value)
            if not order_id or not sku:
                continue
            warehouse = normalize(ws.cell(row, cols["收货仓"]).value)
            if warehouse not in sites:
                raise KeyError(f"仓库信息表未映射收货仓：{warehouse}")
            site = sites[warehouse]
            ship, arrival, normalized_arrival = tracking.get(order_id, (None, None, None))
            rows.append({
                "order": order_id,
                "warehouse": warehouse,
                "site": site,
                "site_batch": station_batch(site, warehouse),
                "sku": sku,
                "qty": ws.cell(row, cols["预计数量"]).value,
                "ship": ship,
                "arrival": arrival,
                "normalized_arrival": normalized_arrival,
            })
        return rows
    finally:
        wb.close()


def normalized_arrival(ship, site) :
    if isinstance(ship, (int, float)):
        ship = from_excel(ship)
    if not isinstance(ship, (date, datetime)):
        return None, None
    days = {"菲律宾": 50, "马来": 40, "印尼": 70, "泰国": 35}[site]
    arrival = ship + timedelta(days=days)
    day = 8 if arrival.day <= 10 else 18 if arrival.day <= 20 else 28
    normalized = date(arrival.year, arrival.month, day)
    return arrival, normalized


def load_unshipped(path) :
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        rows = []
        missing = [name for name in UNSHIPPED_SOURCE_SHEETS if name not in wb.sheetnames]
        if missing:
            raise KeyError(f"A未发数据源缺少站点 Sheet：{missing}")
        for sheet_name in UNSHIPPED_SOURCE_SHEETS:
            ws = wb[sheet_name]
            cols = require_columns(ws, ["SKU", "数量", "要求发货时间", "仓库", "站点"])
            sheet_count = 0
            for row in range(2, ws.max_row + 1):
                sku = normalize(ws.cell(row, cols["SKU"]).value)
                if not sku:
                    continue
                warehouse = normalize(ws.cell(row, cols["仓库"]).value)
                site = normalize(ws.cell(row, cols["站点"]).value)
                if site == "菲律宾" and sku.startswith("HSTEST-"):
                    sku = sku[len("HSTEST-"):]
                ship = ws.cell(row, cols["要求发货时间"]).value
                arrival, normalized_date = normalized_arrival(ship, site)
                rows.append({
                    "order": "A未发",
                    "warehouse": warehouse,
                    "site": site,
                    "site_batch": station_batch(site, warehouse),
                    "sku": sku,
                    "qty": ws.cell(row, cols["数量"]).value,
                    "ship": ship,
                    "arrival": arrival,
                    "normalized_arrival": normalized_date,
                })
                sheet_count += 1
            log(f"A未发源 Sheet 完成：{sheet_name}，有效SKU行={sheet_count}")
        log(f"A未发站点 Sheet 拼接完成：Sheet数={len(UNSHIPPED_SOURCE_SHEETS)}，总行数={len(rows)}；未读取未发数据 Sheet")
        return rows
    finally:
        wb.close()


def sync_mapping_sheets(wb, mapping_path) :
    for name in ("SKU信息表", "店铺信息表", "仓库信息表"):
        count = copy_source_sheet(mapping_path, name, wb[name])
        log(f"Sheet 完成：{name}，数据行={count}")


def prepare_indonesia_sheets(wb) :
    old_name = "SKU批次表_印尼"
    first_name = "SKU批次表_印尼1"
    second_name = "SKU批次表_印尼2"
    if first_name not in wb.sheetnames:
        if old_name in wb.sheetnames:
            wb[old_name].title = first_name
        else:
            wb.create_sheet(first_name)
    elif old_name in wb.sheetnames:
        del wb[old_name]
    if second_name not in wb.sheetnames:
        copied = wb.copy_worksheet(wb[first_name])
        copied.title = second_name


def batch_formula(row, site_batch, batch_index, output_kind, last_order_row) :
    output_expr = {
        "order": 'TEXTJOIN(" / ",TRUE,UNIQUE(FILTER(orderCol,matchCond)))',
        "qty": "SUM(FILTER(qtyCol,matchCond))",
        "ship": "thisShip",
        "arrival": "MIN(FILTER(arrCol,matchCond))",
    }[output_kind]
    end_row = max(last_order_row, 2)
    return (
        f'=IF($A{row}="","",IFERROR(LET(sku,$A{row},site,"{site_batch}",'
        f'orderCol,\'入库单表\'!$A$2:$A${end_row},siteCol,\'入库单表\'!$D$2:$D${end_row},'
        f'skuCol,\'入库单表\'!$E$2:$E${end_row},qtyCol,\'入库单表\'!$G$2:$G${end_row},'
        f'shipCol,\'入库单表\'!$H$2:$H${end_row},arrCol,\'入库单表\'!$J$2:$J${end_row},'
        'baseCond,IFERROR((siteCol=site)*(skuCol=sku)*(shipCol<>""),0),'
        'sortedShip,UNIQUE(SORTBY(FILTER(shipCol,baseCond),FILTER(arrCol,baseCond),1,FILTER(shipCol,baseCond),1)),'
        f'thisShip,INDEX(sortedShip,{batch_index}),matchCond,IFERROR((siteCol=site)*(skuCol=sku)*(shipCol=thisShip),0),'
        f'{output_expr}),""))'
    )


def populate_batch_sheet(ws, site_batch, skus, last_order_row) :
    reset_sheet(ws, BATCH_HEADERS)
    for index, sku in enumerate(skus, 2):
        if index > 2:
            copy_row_style(ws, 2, index, len(BATCH_HEADERS))
        ws.cell(index, 1).value = sku
        style_data_cell(ws.cell(index, 1))
        col = 2
        for batch_no in range(1, 5):
            for kind in ("order", "qty", "ship", "arrival"):
                ws.cell(index, col).value = batch_formula(index, site_batch, batch_no, kind, last_order_row)
                style_data_cell(ws.cell(index, col))
                ws.cell(index, col).fill = FORMULA_FILL
                col += 1
    apply_short_date_format(ws, (4, 5, 8, 9, 12, 13, 16, 17))
    ws.freeze_panes = "A2"
    return len(skus)


def report_over_four_batches(rows) :
    grouped = {}
    for item in rows:
        if item["ship"] in (None, ""):
            continue
        grouped.setdefault((item["site_batch"], item["sku"]), set()).add(item["ship"])
    for site_batch in SITE_BATCHES:
        over = sorted((sku, len(dates)) for (site, sku), dates in grouped.items() if site == site_batch and len(dates) > 4)
        log(f"四批限制校验：{site_batch}，超过四批SKU数={len(over)}，样例={over[:5]}")


def unique_skus_by_site_batch(rows) :
    values = {site_batch: [] for site_batch in SITE_BATCHES}
    seen = {site_batch: set() for site_batch in SITE_BATCHES}
    for item in rows:
        site_batch = normalize(item["site_batch"])
        sku = normalize(item["sku"])
        if site_batch not in values:
            raise ValueError(f"入库单表存在未配置的站点_批次：{site_batch}")
        if sku and sku not in seen[site_batch]:
            seen[site_batch].add(sku)
            values[site_batch].append(sku)
    return values


def write_workbook(paths, month) :
    target = paths["target"]
    wb = load_or_create_workbook(target)
    try:
        prepare_indonesia_sheets(wb)
        ensure_workbook_structure(wb)
        clear_refresh_sheets(wb)
        sync_mapping_sheets(wb, paths["mapping"])
        stock_rows = validate_stock_source(paths["stock"])
        log(f"月初库存快照校验完成：数据行={stock_rows}；该文件供后续分站点净需求阶段使用")
        site_map = warehouse_site_map(paths["mapping"])
        tracking_rows, tracking_lookup = load_tracking(paths["tracking"])
        new_rows = load_new_inbound(paths["in_new"], site_map, tracking_lookup)
        unshipped_rows = load_unshipped(paths["unshipped"])

        tracking_ws = wb["飞书-入库单跟进表"]
        reset_sheet(tracking_ws, ["仓库", "站点", "入库单号", "发货时间", "到仓时间", "到仓时间整合", "是否到货", "到货天数"])
        append_rows(tracking_ws, tracking_rows)
        for row in range(2, tracking_ws.max_row + 1):
            tracking_ws.cell(row, 8).value = f'=IFERROR(E{row}-D{row},"")'
            tracking_ws.cell(row, 8).fill = FORMULA_FILL
        apply_short_date_format(tracking_ws, (4, 5, 6))
        log(f"Sheet 完成：飞书-入库单跟进表，数据行={len(tracking_rows)}，公式={max(tracking_ws.max_row - 1, 0)}")

        new_ws = wb["千易入库单_新建"]
        reset_sheet(new_ws, ["入库单号", "收货仓", "站点（中文）", "商品编码", "预计数量", "发货时间", "到仓时间", "到仓时间整合"])
        for index, item in enumerate(new_rows, 2):
            if index > 2:
                copy_row_style(new_ws, 2, index, 8)
            values = [item["order"], item["warehouse"], f'=VLOOKUP(B{index},\'仓库信息表\'!A:C,2,0)', item["sku"], item["qty"],
                      f'=IFERROR(VLOOKUP(A{index},\'飞书-入库单跟进表\'!C:H,2,0),"")',
                      f'=IFERROR(VLOOKUP(A{index},\'飞书-入库单跟进表\'!C:H,3,0),"")',
                      f'=IFERROR(VLOOKUP(A{index},\'飞书-入库单跟进表\'!C:H,4,0),"")']
            for col, value in enumerate(values, 1):
                new_ws.cell(index, col).value = value
                style_data_cell(new_ws.cell(index, col))
                if col in (3, 6, 7, 8):
                    new_ws.cell(index, col).fill = FORMULA_FILL
        apply_short_date_format(new_ws, (6, 7, 8))
        log(f"Sheet 完成：千易入库单_新建，数据行={len(new_rows)}，公式={len(new_rows) * 4}")

        a_ws = wb["A未发数据表"]
        reset_sheet(a_ws, ["站点", "仓库", "SKU", "数量", "要求发货时间", "到仓时间", "到仓时间整合"])
        append_rows(a_ws, [[x["site"], x["warehouse"], x["sku"], x["qty"], x["ship"], x["arrival"], x["normalized_arrival"]] for x in unshipped_rows])
        apply_short_date_format(a_ws, (5, 6, 7))
        log(f"Sheet 完成：A未发数据表，数据行={len(unshipped_rows)}，菲律宾 HSTEST- 前缀已清理")

        order_ws = wb["入库单表"]
        reset_sheet(order_ws, ["入库单号", "站点（中文）", "仓库", "站点_批次", "商品编码", "预计数量", "预计数量Value", "发货时间", "到仓时间", "到仓时间整合"])
        all_rows = new_rows + unshipped_rows
        for index, item in enumerate(all_rows, 2):
            if index > 2:
                copy_row_style(order_ws, 2, index, 10)
            values = [item["order"], item["site"], item["warehouse"], item["site_batch"], item["sku"], item["qty"],
                      f'=IFERROR(F{index}+0,0)', item["ship"], item["arrival"], item["normalized_arrival"]]
            for col, value in enumerate(values, 1):
                order_ws.cell(index, col).value = value
                style_data_cell(order_ws.cell(index, col))
                if col == 7:
                    order_ws.cell(index, col).fill = FORMULA_FILL
        apply_short_date_format(order_ws, (8, 9, 10))
        counts = {site: sum(1 for row in all_rows if row["site_batch"] == site) for site in SITE_BATCHES}
        log(f"Sheet 完成：入库单表，数据行={len(all_rows)}，站点_批次分布={counts}")
        report_over_four_batches(all_rows)

        site_skus = unique_skus_by_site_batch(all_rows)
        for site_batch in SITE_BATCHES:
            name = f"SKU批次表_{site_batch}"
            row_count = populate_batch_sheet(wb[name], site_batch, site_skus[site_batch], order_ws.max_row)
            log(f"Sheet 完成：{name}，SKU行={row_count}，公式={row_count * 16}")

        missing_sheets = set(FINAL_SHEET_ORDER) - set(wb.sheetnames)
        if missing_sheets:
            raise KeyError(f"在途批次表缺少 Sheet：{sorted(missing_sheets)}")
        wb._sheets = [wb[name] for name in FINAL_SHEET_ORDER]

        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
        temp_dir = paths["target"].parent.parent / ".codex_tmp"
        temp_dir.mkdir(exist_ok=True)
        temp_path = temp_dir / f"{target.stem}_building.xlsx"
        wb.save(temp_path)
    finally:
        wb.close()
    shutil.move(str(temp_path), str(target))
    log(f"工作簿已写入：{target}")


def recalculate_with_excel(path, timeout_seconds = 900) :
    log("开始调用 Microsoft Excel 全量计算并刷新缓存")
    script = r'''
param([string]$WorkbookPath, [int]$TimeoutSeconds)
$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.AskToUpdateLinks = $false
    $book = $excel.Workbooks.Open($WorkbookPath, 0, $false)
    $excel.CalculateFullRebuild()
    $start = Get-Date
    while ($excel.CalculationState -ne 0) {
        if (((Get-Date) - $start).TotalSeconds -gt $TimeoutSeconds) { throw 'Excel 计算超时' }
        Start-Sleep -Milliseconds 500
    }
    $book.Save()
}
finally {
    if ($book -ne $null) { $book.Close($true) | Out-Null; [void][Runtime.InteropServices.Marshal]::ReleaseComObject($book) }
    if ($excel -ne $null) { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel) }
    [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}
'''
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as handle:
            handle.write(script)
            script_path = Path(handle.name)
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script_path),
             "-WorkbookPath", str(path.resolve()), "-TimeoutSeconds", str(timeout_seconds)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout_seconds + 60,
        )
        if completed.returncode != 0:
            raise RuntimeError("Excel 缓存刷新失败：\n" + completed.stderr.strip())
    finally:
        if script_path:
            try:
                script_path.unlink()
            except FileNotFoundError:
                pass
    log("Microsoft Excel 全量计算、保存和关闭完成")


def validate_result(path) :
    formula_wb = load_workbook(path, read_only=False, data_only=False)
    value_wb = load_workbook(path, read_only=False, data_only=True)
    try:
        expected = {f"SKU批次表_{site}" for site in SITE_BATCHES}
        missing = expected - set(formula_wb.sheetnames)
        if missing:
            raise ValueError(f"结果缺少批次 Sheet：{sorted(missing)}")
        if "SKU批次表_印尼" in formula_wb.sheetnames:
            raise ValueError("旧 Sheet SKU批次表_印尼 仍然存在")
        short_date_columns = {
            "飞书-入库单跟进表": (4, 5, 6),
            "千易入库单_新建": (6, 7, 8),
            "A未发数据表": (5, 6, 7),
            "入库单表": (8, 9, 10),
        }
        for site_batch in SITE_BATCHES:
            short_date_columns[f"SKU批次表_{site_batch}"] = (4, 5, 8, 9, 12, 13, 16, 17)
        date_format_errors = []
        for sheet_name, columns in short_date_columns.items():
            ws = formula_wb[sheet_name]
            for col in columns:
                for row in range(2, ws.max_row + 1):
                    if not is_short_date_format(ws.cell(row, col).number_format):
                        date_format_errors.append(f"{sheet_name}!{ws.cell(row, col).coordinate}")
                        break
        if date_format_errors:
            raise ValueError("以下日期列未使用短日期格式：" + "、".join(date_format_errors))
        order_ws = value_wb["入库单表"]
        cols = require_columns(order_ws, ["站点（中文）", "仓库", "站点_批次", "商品编码"])
        errors = []
        expected_skus = {site_batch: [] for site_batch in SITE_BATCHES}
        seen_skus = {site_batch: set() for site_batch in SITE_BATCHES}
        for row in range(2, order_ws.max_row + 1):
            site = normalize(order_ws.cell(row, cols["站点（中文）"]).value)
            warehouse = normalize(order_ws.cell(row, cols["仓库"]).value)
            actual = normalize(order_ws.cell(row, cols["站点_批次"]).value)
            sku = normalize(order_ws.cell(row, cols["商品编码"]).value)
            if site == "印尼" and actual not in {"印尼1", "印尼2"}:
                errors.append((row, warehouse, actual))
            if actual in expected_skus and sku and sku not in seen_skus[actual]:
                seen_skus[actual].add(sku)
                expected_skus[actual].append(sku)
        if errors:
            raise ValueError(f"仍有印尼记录未正确拆分，示例：{errors[:5]}")
        grain_errors = []
        for site_batch in SITE_BATCHES:
            batch_ws = value_wb[f"SKU批次表_{site_batch}"]
            actual_skus = [normalize(batch_ws.cell(row, 1).value) for row in range(2, batch_ws.max_row + 1)]
            actual_skus = [sku for sku in actual_skus if sku]
            if actual_skus != expected_skus[site_batch]:
                grain_errors.append(
                    f"{site_batch}：预期{len(expected_skus[site_batch])}行，实际{len(actual_skus)}行"
                )
        if grain_errors:
            raise ValueError("SKU批次表粒度或排序与入库单表不一致：" + "；".join(grain_errors))
        formula_errors = []
        for ws in value_wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?"}:
                        formula_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
                        if len(formula_errors) >= 20:
                            break
                if len(formula_errors) >= 20:
                    break
            if len(formula_errors) >= 20:
                break
        if formula_errors:
            raise ValueError("发现公式错误：\n" + "\n".join(formula_errors))
        log(f"校验完成：Sheet={len(formula_wb.sheetnames)}，入库单表数据行={order_ws.max_row - 1}，未发现关键公式错误")
    finally:
        formula_wb.close()
        value_wb.close()


def run(root, month_text) :
    month = parse_month(month_text)
    paths = resolve_paths(root.resolve(), month)
    log(f"根目录：{root.resolve()}")
    log(f"执行月份：{month_text}")
    log(f"在途批次日快照：{next_month_first_day(month):%Y-%m-%d}；不检索或回退到其他日期")
    for name in ("in_new", "tracking", "stock", "unshipped", "mapping", "target"):
        log(f"输入确认 {name}：{paths[name]}")
    try:
        write_workbook(paths, month)
        recalculate_with_excel(paths["target"])
        validate_result(paths["target"])
        temp_dir = root / ".codex_tmp"
        if temp_dir.exists() and not any(temp_dir.iterdir()):
            temp_dir.rmdir()
    finally:
        release_resources_and_verify([paths["target"]])
    return paths["target"]


def main() :
    parser = argparse.ArgumentParser(description="生成并刷新月度在途批次表")
    parser.add_argument("--root", type=Path, default=ROOT_DIR, help="项目根目录")
    parser.add_argument("--month", default=DEFAULT_MONTH, help="月度日期，格式 YYYY-MM")
    args = parser.parse_args()
    output = run(args.root, args.month)
    log(f"执行成功：{output}")


if __name__ == "__main__":
    main()
